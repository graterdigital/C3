#!/usr/bin/env python3
"""
Client update & change-request document — EXPORT and IMPORT.

  bin/copy-doc.py export [out.docx]                 → Word doc: every piece of site copy (ID | Current | Your change),
                                                      every image slot (thumbnail | where used | paste a link), events,
                                                      partners, plus a free-form change-request table.
  bin/copy-doc.py import edited.docx --dry          → report what the client changed (copy, images, events, partners, requests)
  bin/copy-doc.py import edited.docx [--images DIR] → write changes to content/*.json + config/site.php; download/crop images
                                                      (DIR = folder of already-downloaded files named <slot>.<ext> for links the
                                                      sandbox can't fetch). Then run the seed.

Sources / targets:
  content/copy.json      pattern copy (hero, sections, about, donate, FAQ, forms)   keys: copy.<key>
  content/site.json      pages (title, excerpt, blocks), events, posts             keys: page.<slug>.*, event.<slug>.*, post.<slug>.*
  config/site.php        name, tagline, email, address, socials, CTA, donate       keys: site.<key>
  config/site.php        partners[]                                                keys: partner.<n>.name/url  (+ image slot partner-*.png)
  content/images.json    image slots                                               keys: img.<file>

Round-trip contract: column 1 (ID) is the key. The client writes only in the "Your change" column; blank = keep.
"""
import json, re, sys, io, pathlib, hashlib, urllib.request
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = pathlib.Path(__file__).resolve().parents[1]
C = ROOT / 'content'
IMG = ROOT / 'wp-content/themes/gd-starter/assets/img'
CFG = ROOT / 'wp-content/themes/gd-starter/config/site.php'
INK, RED, NAVY, MUTED = RGBColor(0x1c, 0x1f, 0x24), RGBColor(0xd3, 0x2f, 0x2a), RGBColor(0x0f, 0x1e, 0x38), RGBColor(0x4a, 0x52, 0x60)
SITE_NAME = 'C3 · Creative Collaboration for Compassion'
SITE_URL = 'https://c3collaborative.kinsta.cloud'

def J(name):
    return json.load(open(C / name, encoding='utf-8'))

# ----------------------------------------------------------------------------- config/site.php
CFG_KEYS = [('name', 'Site name'), ('tagline', 'Tagline (footer)'), ('email', 'Contact email'), ('address', 'Address line (footer / contact)'),
            ('phone', 'Phone (shown if set)'), ('instagram', 'Instagram URL'), ('facebook', 'Facebook URL'),
            ('cta_label', 'Header button label'), ('donate_url', 'Donate link (GoFundMe)'), ('donate_label', 'Donate button label'),
            ('form_recipient', 'Form notifications go to (blank = contact email)')]

def cfg_src():
    return open(CFG, encoding='utf-8').read()

def cfg_value(key):
    m = re.search(r"'%s'\s*=>\s*'((?:[^'\\]|\\.)*)'" % re.escape(key), cfg_src())
    return m.group(1).replace("\\'", "'") if m else ''

def cfg_set(src, key, value):
    v = value.replace("'", "\\'")
    return re.sub(r"('%s'\s*=>\s*)'(?:[^'\\]|\\.)*'" % re.escape(key), lambda m: m.group(1) + "'" + v + "'", src, count=1)

def cfg_partners():
    m = re.search(r"'partners'\s*=>\s*\[(.*?)\n\t\],", cfg_src(), re.S)
    out = []
    for row in re.findall(r"\[\s*'name'\s*=>\s*'((?:[^'\\]|\\.)*)',\s*'logo'\s*=>\s*'([^']*)',\s*'url'\s*=>\s*'([^']*)'\s*\]", m.group(1)):
        out.append({'name': row[0].replace("\\'", "'"), 'logo': row[1], 'url': row[2]})
    return out

def cfg_write_partners(src, partners):
    body = ''.join("\t\t[ 'name' => '%s', 'logo' => '%s', 'url' => '%s' ],\n" % (p['name'].replace("'", "\\'"), p['logo'], p['url']) for p in partners)
    return re.sub(r"('partners'\s*=>\s*\[\n)(.*?)(\n\t\],)", lambda m: m.group(1) + body.rstrip('\n') + m.group(3), src, count=1, flags=re.S)

# ----------------------------------------------------------------------------- rows
def build_rows():
    """Yield (section_title, note, kind, rows). kind: copy | image | request. copy rows: (id, label, current, multiline)."""
    copy = J('copy.json'); site = J('site.json'); images = J('images.json')

    rows = [(f'site.{k}', lab, cfg_value(k), False) for k, lab in CFG_KEYS]
    rows.insert(1, ('site.blogdescription', 'Site description (browser tab / search result)', site['site'].get('tagline', ''), False))
    yield ('Site-wide details', 'Name, contact details, links and buttons that appear on every page.', 'copy', rows)

    groups = [('hero', 'Home — hero'), ('fixlist', 'Home — yellow strip (what we fix, one item per row)'), ('intro', 'Home — Neighbors with tools'),
              ('services', 'Home — What we fix'), ('steps', 'Home — How it works'), ('band', 'Home — photo band caption'),
              ('doors', 'Home — Request help / Volunteer blocks'), ('events', 'Home — Events section'), ('partners', 'Home — Partners section'),
              ('instagram', 'Home — Instagram section'), ('footer', 'Every page — bottom call-to-action band'),
              ('mission', 'About — mission block'), ('about', 'About — story'), ('loop', 'About — How the free repairs get paid for'),
              ('donate', 'Donate page'), ('involved', 'Get involved page — intros and thank-you messages'), ('contact', 'Contact page'), ('faq', 'FAQ page')]
    for prefix, title in groups:
        keys = [k for k in copy if k.split('.')[0] == prefix]
        rows = [(f'copy.{k}', k.split('.', 1)[1].replace('.', ' › '), copy[k], len(copy[k]) > 120) for k in keys]
        if prefix == 'faq':
            n = len([k for k in keys if k.startswith('faq.q')])
            for j in range(n + 1, n + 3):
                rows.append((f'copy.faq.q{j}', f'q{j} (new question — optional)', '', False))
                rows.append((f'copy.faq.a{j}', f'a{j} (its answer)', '', True))
        if prefix == 'fixlist':
            for j in range(len(keys) + 1, len(keys) + 3):
                rows.append((f'copy.fixlist.{j}', f'{j} (optional extra item)', '', False))
        yield (title, '', 'copy', rows)

    # pages (titles/excerpts/blocks that live in site.json)
    for pg in site['pages']:
        if pg['slug'] == 'home':
            continue
        rows = [(f"page.{pg['slug']}.title", 'Page title', pg['title'], False), (f"page.{pg['slug']}.excerpt", 'Subtitle under the title', pg.get('excerpt', ''), False)]
        for i, b in enumerate(pg.get('blocks', [])):
            if 'h2' in b: rows.append((f"page.{pg['slug']}.b{i}.h2", 'Heading', b['h2'], False))
            elif 'p' in b: rows.append((f"page.{pg['slug']}.b{i}.p", 'Paragraph', b['p'], True))
            elif 'list' in b: rows.append((f"page.{pg['slug']}.b{i}.list", 'List (one item per line)', '\n'.join(b['list']), True))
        if len(rows) > 2 or pg['slug'] in ('what-we-do', 'updates'):
            yield (f"Page: {pg['title']}  ({SITE_URL}/{pg['slug']}/)", '', 'copy', rows)

    # partners
    rows = []
    for i, p in enumerate(cfg_partners(), 1):
        rows.append((f'partner.{i}.name', f'Partner {i} — name', p['name'], False))
        rows.append((f'partner.{i}.url', f'Partner {i} — website or Instagram link', p['url'], False))
    for i in range(len(cfg_partners()) + 1, len(cfg_partners()) + 4):
        rows.append((f'partner.{i}.name', f'Partner {i} — name (new, optional)', '', False))
        rows.append((f'partner.{i}.url', f'Partner {i} — link', '', False))
    yield ('Partners & affiliates', 'Shown on the home page and About page. Logos are in the Images section below.', 'copy', rows)

    # events
    rows = []
    for ev in site.get('events', []):
        s = ev['slug']
        for k, lab, ml in [('title', 'Title', False), ('excerpt', 'One-line summary', False), ('content', 'Description', True),
                           ('start', 'Starts (YYYY-MM-DD HH:MM)', False), ('end', 'Ends', False), ('venue', 'Venue', False),
                           ('address', 'Address', False), ('cost', 'Cost', False), ('organizer', 'Hosted with', False),
                           ('link', 'RSVP / info link', False), ('status', 'Status (publish / draft)', False)]:
            rows.append((f'event.{s}.{k}', lab, str(ev.get(k, 'publish' if k == 'status' else '')), ml))
    for n in (1, 2):
        for k, lab, ml in [('title', 'Title', False), ('excerpt', 'One-line summary', False), ('content', 'Description', True),
                           ('start', 'Starts (YYYY-MM-DD HH:MM)', False), ('end', 'Ends', False), ('venue', 'Venue', False),
                           ('address', 'Address', False), ('cost', 'Cost', False), ('link', 'RSVP / info link', False)]:
            rows.append((f'event.new{n}.{k}', f'New event {n} — {lab}', '', ml))
    yield ('Events', 'Existing events, plus two blank slots for new ones. Posters go in the Images section (or paste a link in the row notes).', 'copy', rows)

    # posts
    rows = []
    for po in site.get('posts', []):
        rows.append((f"post.{po['slug']}.title", 'Title', po['title'], False))
        rows.append((f"post.{po['slug']}.excerpt", 'Summary', po.get('excerpt', ''), False))
        rows.append((f"post.{po['slug']}.content", 'Body (one paragraph per line)', po.get('content', ''), True))
    rows.append(('post.new1.title', 'New update — title', '', False))
    rows.append(('post.new1.content', 'New update — body', '', True))
    yield ('Updates (blog)', '', 'copy', rows)

    # images
    rows = []
    for file, m in images.items():
        if file.startswith('_'):
            continue
        rows.append((f'img.{file}', m, (IMG / file) if (IMG / file).exists() else None, False))
    yield ('Images', 'Paste a link to the replacement photo in the right-hand column (Google Drive, Dropbox, iCloud — set sharing to “anyone with the link”). We crop and resize to fit the slot.', 'image', rows)

    yield ('Other change requests', 'Anything not covered above: a page you want added, a section moved, a link changed, something that looks wrong on your phone.', 'request', [])

# ----------------------------------------------------------------------------- docx helpers
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)

def set_cell(cell, text, size=9.5, bold=False, color=None, italic=False):
    cell.text = ''
    paras = text.split('\n') if text else ['']
    for i, line in enumerate(paras):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if color: r.font.color.rgb = color

def header_row(tbl, labels, widths):
    for c, w in zip(tbl.rows[0].cells, widths): c.width = w
    for c, t in zip(tbl.rows[0].cells, labels):
        set_cell(c, t, 9, True, RGBColor(0xff, 0xff, 0xff)); shade(c, '0F1E38')
    trPr = tbl.rows[0]._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)

def thumb(path, max_px=360):
    from PIL import Image
    im = Image.open(path); im.thumbnail((max_px, max_px))
    if im.mode in ('RGBA', 'P'):
        bg = Image.new('RGB', im.size, (240, 240, 240)); bg.paste(im.convert('RGBA'), mask=im.convert('RGBA').split()[-1]); im = bg
    buf = io.BytesIO(); im.convert('RGB').save(buf, 'JPEG', quality=70); buf.seek(0)
    return buf, im.size

# ----------------------------------------------------------------------------- export
def export(out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(1.5); sec.top_margin = sec.bottom_margin = Cm(1.5)
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10)

    h = doc.add_paragraph(); r = h.add_run(SITE_NAME); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = INK
    p = doc.add_paragraph(); r = p.add_run('Website update & change request document'); r.font.size = Pt(13); r.font.color.rgb = RED
    p = doc.add_paragraph(); r = p.add_run(f'Preview: {SITE_URL}   ·   Prepared by Grater Digital'); r.font.size = Pt(9.5); r.font.color.rgb = MUTED
    doc.add_paragraph()
    for t, b in [
        ('How to use this document', True),
        ('1. Every piece of text on the site is listed below with its current wording. To change something, write the new text in the “Your change” column. Leave it blank to keep what’s there.', False),
        ('2. Don’t edit the ID column or delete rows — the ID is how we match your change to the right spot on the site.', False),
        ('3. Lists (the yellow strip, FAQ) have a few blank rows at the end for additions. To remove an item, write “REMOVE” in its Your change column.', False),
        ('4. Images: each photo slot on the site is shown with a thumbnail. To swap one, paste a link to the new photo in the right-hand column (Google Drive, Dropbox or iCloud link set to “anyone with the link”). We crop and resize it to fit. The slot notes tell you which orientation works best.', False),
        ('5. Events: update dates, venues and links in place; two blank slots are provided for new events. Set Status to “draft” to hide an event or “publish” to show it.', False),
        ('6. Anything else — a page you want, a section you want moved, something that looks wrong on your phone — goes in the “Other change requests” table at the end.', False),
        ('7. Send the file back as Word (.docx). We’ll apply everything in one update and send you a before/after.', False),
    ]:
        p = doc.add_paragraph(); r = p.add_run(t); r.bold = b; r.font.size = Pt(10.5 if b else 10)
    doc.add_page_break()

    sections = list(build_rows())
    p = doc.add_paragraph(); r = p.add_run('Contents'); r.bold = True; r.font.size = Pt(14)
    for title, _, _, _ in sections:
        doc.add_paragraph(title, style='List Number')
    doc.add_page_break()

    n_fields = 0
    for si, (title, note, kind, rows) in enumerate(sections):
        if si and kind != 'copy' or (si and sections[si-1][2] != 'copy'):
            doc.add_page_break()
        hp = doc.add_paragraph(); hr = hp.add_run(title); hr.bold = True; hr.font.size = Pt(14); hr.font.color.rgb = INK
        hp.paragraph_format.keep_with_next = True
        if note:
            np_ = doc.add_paragraph(); nr = np_.add_run(note); nr.italic = True; nr.font.size = Pt(9); nr.font.color.rgb = MUTED

        if kind == 'copy':
            tbl = doc.add_table(rows=1, cols=3); tbl.style = 'Table Grid'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            widths = [Cm(3.6), Cm(7.5), Cm(7.5)]
            header_row(tbl, ['ID', 'Current text', 'Your change (blank = keep)'], widths)
            for rid, label, cur, ml in rows:
                cells = tbl.add_row().cells
                for c, w in zip(cells, widths): c.width = w
                set_cell(cells[0], rid + '\n' + label, 7.5, False, MUTED); shade(cells[0], 'F3F0E8')
                set_cell(cells[1], cur, 9.5, color=RGBColor(0x44, 0x4a, 0x55))
                set_cell(cells[2], '', 9.5)
                n_fields += 1
        elif kind == 'image':
            tbl = doc.add_table(rows=1, cols=3); tbl.style = 'Table Grid'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            widths = [Cm(5.0), Cm(6.6), Cm(7.0)]
            header_row(tbl, ['Current image', 'Where it’s used / what fits', 'Link to replacement (paste here)'], widths)
            for rid, m, path, _ in rows:
                cells = tbl.add_row().cells
                for c, w in zip(cells, widths): c.width = w
                cells[0].text = ''
                if path:
                    buf, (w, hgt) = thumb(path)
                    run = cells[0].paragraphs[0].add_run(); run.add_picture(buf, width=Cm(4.6) if w >= hgt else Cm(3.2))
                else:
                    set_cell(cells[0], '(no image yet)', 9, italic=True, color=MUTED)
                shape = 'Square' if abs(m['ratio'] - 1) < .05 else ('Portrait' if m['ratio'] < 1 else ('Wide landscape' if m['ratio'] > 2 else 'Landscape'))
                tip = m.get('tip', '')
                note = tip if tip.lower().startswith(('square', 'portrait', 'landscape', 'wide')) else f'{shape}. {tip}'
                set_cell(cells[1], f"{rid}\n{m['used']}\n\n{note}", 8.5)
                set_cell(cells[2], '', 9.5)
                n_fields += 1
        else:
            tbl = doc.add_table(rows=1, cols=3); tbl.style = 'Table Grid'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            widths = [Cm(4.2), Cm(10.8), Cm(3.6)]
            header_row(tbl, ['Page / section', 'What you’d like changed', 'Priority (must / nice)'], widths)
            for i in range(8):
                row = tbl.add_row(); row.height = Cm(1.8); cells = row.cells
                for c, w in zip(cells, widths): c.width = w
                set_cell(cells[0], f'request.{i+1}', 7.5, color=MUTED); shade(cells[0], 'F3F0E8')
                for c in cells[1:]:
                    set_cell(c, '', 9.5)
        doc.add_paragraph()
    doc.save(out)
    print(f'wrote {out}: {len(sections)} sections, {n_fields} fields')

# ----------------------------------------------------------------------------- import
def cell_text(cell):
    return '\n'.join(p.text for p in cell.paragraphs).strip()

def read_doc(path):
    doc = Document(path)
    copy, images, requests = {}, {}, []
    for t in doc.tables:
        for row in t.rows[1:]:
            cells = row.cells
            if len(cells) < 3:
                continue
            first = cells[0].text.strip().split('\n')[0].strip()
            mid_first = cells[1].text.strip().split('\n')[0].strip()
            if first.startswith('request.'):
                if cell_text(cells[1]):
                    requests.append((cell_text(cells[1]), cell_text(cells[2])))
            elif mid_first.startswith('img.'):
                v = cell_text(cells[2])
                if v:
                    images[mid_first[4:]] = v
            elif re.match(r'^(site|copy|page|partner|event|post)\.', first):
                v = cell_text(cells[2])
                if v:
                    copy[first] = v
    return copy, images, requests

def fetch(url, dest):
    """Download a share link. Google Drive / Dropbox links are rewritten to direct-download form."""
    m = re.search(r'drive\.google\.com/(?:file/d/|open\?id=|uc\?id=)([\w-]+)', url)
    if m:
        url = f'https://drive.google.com/uc?export=download&id={m.group(1)}'
    if 'dropbox.com' in url:
        url = re.sub(r'[?&]dl=0', '', url) + ('&' if '?' in url else '?') + 'dl=1'
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=60) as r:
        data = r.read()
    dest.write_bytes(data)
    return dest

def place_image(slot, src, meta):
    from PIL import Image, ImageOps
    im = ImageOps.exif_transpose(Image.open(src))
    if meta['kind'] == 'logo':
        im = im.convert('RGBA'); im = im.crop(im.getbbox() or (0, 0, im.width, im.height)); im.thumbnail((meta['width'], meta['width']))
        im.save(IMG / slot, optimize=True)
    else:
        im = im.convert('RGB'); W, H = im.size; r = meta['ratio']
        if W / H > r:
            nw = int(H * r); x = (W - nw) // 2; im = im.crop((x, 0, x + nw, H))
        else:
            nh = int(W / r); y = int((H - nh) * 0.4); im = im.crop((0, y, W, y + nh))
        if im.width > meta['width']:
            im = im.resize((meta['width'], int(im.height * meta['width'] / im.width)), Image.LANCZOS)
        im.save(IMG / slot, 'JPEG', quality=82, optimize=True, progressive=True)
    return im.size

def import_doc(path, dry=False, images_dir=None):
    copy_changes, image_links, requests = read_doc(path)
    site = J('site.json'); copy = J('copy.json'); images = J('images.json')
    print(f'{len(copy_changes)} text changes, {len(image_links)} image links, {len(requests)} change requests')
    for rid, v in copy_changes.items():
        print(f'  {rid}: {v[:90]!r}')
    for slot, url in image_links.items():
        print(f'  img {slot} ← {url}')
    for what, prio in requests:
        print(f'  REQUEST [{prio or "-"}]: {what}')
    if dry:
        return

    src = cfg_src(); partners = cfg_partners(); ev_by = {e['slug']: e for e in site.get('events', [])}; po_by = {p['slug']: p for p in site.get('posts', [])}
    new_events, new_posts = {}, {}
    for rid, v in copy_changes.items():
        parts = rid.split('.')
        remove = v.strip().upper() == 'REMOVE'
        if parts[0] == 'copy':
            key = rid[5:]
            if remove: copy.pop(key, None)
            else: copy[key] = v
        elif parts[0] == 'site':
            if parts[1] == 'blogdescription': site['site']['tagline'] = v
            else: src = cfg_set(src, parts[1], v)
        elif parts[0] == 'page':
            pg = next(p for p in site['pages'] if p['slug'] == parts[1])
            if parts[2] in ('title', 'excerpt'): pg[parts[2]] = v
            else:
                i = int(parts[2][1:]); kind = parts[3]
                pg['blocks'][i][kind] = [l for l in v.split('\n') if l.strip()] if kind == 'list' else v
        elif parts[0] == 'partner':
            i = int(parts[1]) - 1
            while len(partners) <= i: partners.append({'name': '', 'logo': '', 'url': ''})
            if remove: partners[i]['name'] = ''
            else: partners[i][parts[2]] = v
        elif parts[0] == 'event':
            slug, k = parts[1], parts[2]
            if slug.startswith('new'):
                new_events.setdefault(slug, {})[k] = v
            elif slug in ev_by:
                ev_by[slug][k] = v
        elif parts[0] == 'post':
            slug, k = parts[1], parts[2]
            if slug.startswith('new'): new_posts.setdefault(slug, {})[k] = v
            elif slug in po_by: po_by[slug][k] = v
    partners = [p for p in partners if p['name']]
    for slug, ev in new_events.items():
        if ev.get('title'):
            ev['slug'] = re.sub(r'[^a-z0-9]+', '-', ev['title'].lower()).strip('-'); ev.setdefault('status', 'draft')
            site.setdefault('events', []).append(ev)
    for slug, po in new_posts.items():
        if po.get('title'):
            po['slug'] = re.sub(r'[^a-z0-9]+', '-', po['title'].lower()).strip('-')
            site.setdefault('posts', []).append(po)

    # images
    pending = []
    for slot, url in image_links.items():
        meta = images.get(slot)
        if not meta:
            print(f'  ! unknown slot {slot}'); continue
        local = None
        if images_dir:
            for cand in pathlib.Path(images_dir).glob(pathlib.Path(slot).stem + '.*'):
                local = cand; break
        if not local:
            try:
                local = fetch(url, pathlib.Path('/tmp') / ('gd-' + slot))
            except Exception as e:
                print(f'  ! could not download {slot}: {e}'); pending.append((slot, url)); continue
        size = place_image(slot, local, meta)
        print(f'  placed {slot} {size}')
        # new partner logo → wire into config
        if slot.startswith('partner-'):
            for p in partners:
                if not p['logo'] and slot.split('-')[1].split('.')[0] in re.sub(r'[^a-z]', '', p['name'].lower()):
                    p['logo'] = slot
    src = cfg_write_partners(src, partners)
    json.dump(copy, open(C / 'copy.json', 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
    json.dump(site, open(C / 'site.json', 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    open(CFG, 'w', encoding='utf-8').write(src)
    if requests:
        (C / 'CHANGE-REQUESTS.md').write_text('# Change requests from client doc\n\n' + ''.join(f'- [ ] ({p or "-"}) {w}\n' for w, p in requests), encoding='utf-8')
        print('  change requests → content/CHANGE-REQUESTS.md')
    if pending:
        print('\nImages that need fetching by hand (save as <slot> into a folder, re-run with --images DIR):')
        for slot, url in pending: print(f'  {slot}  {url}')
    print('content updated — run the seed (images are re-imported when the file changed)')

if __name__ == '__main__':
    if len(sys.argv) < 2 or sys.argv[1] not in ('export', 'import'):
        print(__doc__); sys.exit(1)
    if sys.argv[1] == 'export':
        out = sys.argv[2] if len(sys.argv) > 2 and not sys.argv[2].startswith('--') else str(ROOT / 'preview/C3-Website-Updates.docx')
        pathlib.Path(out).parent.mkdir(parents=True, exist_ok=True)
        export(out)
    else:
        imgs = sys.argv[sys.argv.index('--images') + 1] if '--images' in sys.argv else None
        import_doc(sys.argv[2], '--dry' in sys.argv, imgs)
